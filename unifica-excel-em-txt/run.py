import os

from openpyxl import load_workbook


def get_excel_files(from_directory):
    return [f for f in os.listdir(from_directory) if f.endswith('.xlsx') or f.endswith('.xls')]


def unify_excel_files(from_directory, destination_file, include_header):
    excel_files = get_excel_files(from_directory)

    with open(destination_file, 'w', encoding='utf-8') as out_f:
        for file in excel_files:
            print(f'Copiando {file}')
            file_path = os.path.join(from_directory, file)
            workbook = load_workbook(file_path, data_only=True, read_only=True)

            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)

                if not include_header:
                    next(rows, None)

                lines = [
                    '\t'.join('' if value is None else str(value) for value in row) + '\n'
                    for row in rows
                ]
                out_f.writelines(lines)

            workbook.close()


if __name__ == '__main__':
    from_directory = 'C:\\TEMP G\\'
    destination_file = 'C:\\TEMP G\\FichaFinanceira.txt'
    include_header = 's'

    unify_excel_files(from_directory, destination_file, include_header == 'n')
